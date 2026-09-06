"""애니메이션 단계 펼치기 self-check: python3 python/test_animation.py

LibreOffice 없이 도는 XML 단위 테스트. 애니메이션이 들어간 pptx를 즉석에서 만들어
expand_animation_steps가 슬라이드를 클릭 단계만큼 늘리는지, 각 단계에 보여야 할 것만
남는지 확인한다.
"""
import os
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from ppt_tool import _A_NS, _P_NS, expand_animation_steps

_NS_DECL = f'xmlns:p="{_P_NS}" xmlns:a="{_A_NS}"'


def _effect(node_id, spid, preset="entr", paragraph=None):
    text_range = (
        f'<p:txEl><p:pRg st="{paragraph}" end="{paragraph}"/></p:txEl>'
        if paragraph is not None
        else ""
    )
    return f'''
      <p:par>
        <p:cTn id="{node_id}" presetID="1" presetClass="{preset}" fill="hold"
               grpId="0" nodeType="clickEffect">
          <p:childTnLst>
            <p:set>
              <p:cBhvr>
                <p:cTn id="{node_id + 1}" dur="1" fill="hold"/>
                <p:tgtEl><p:spTgt spid="{spid}">{text_range}</p:spTgt></p:tgtEl>
                <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
              </p:cBhvr>
              <p:to><p:strVal val="visible"/></p:to>
            </p:set>
          </p:childTnLst>
        </p:cTn>
      </p:par>'''


def _timing(clicks):
    groups = "".join(
        f'<p:par><p:cTn id="{500 + index}" fill="hold">'
        f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
        f"<p:childTnLst>{click}</p:childTnLst></p:cTn></p:par>"
        for index, click in enumerate(clicks)
    )
    return f'''<p:timing {_NS_DECL}>
      <p:tnLst>
        <p:par>
          <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
            <p:childTnLst>
              <p:seq concurrent="1" nextAc="seek">
                <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                  <p:childTnLst>{groups}</p:childTnLst>
                </p:cTn>
              </p:seq>
            </p:childTnLst>
          </p:cTn>
        </p:par>
      </p:tnLst>
    </p:timing>'''


def build_sample(path):
    """슬라이드 3장짜리 샘플.

    1장: 도형 3개 — 2·3번이 클릭마다 등장 (도형 단위)
    2장: 문단 3줄짜리 상자 1개 — 2·3번째 줄이 클릭마다 등장 (문단 단위)
    3장: 도형 2개 — 1클릭에 하나 등장, 2클릭에 처음 것이 사라짐 (등장 + 사라짐)
    """
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)

    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape_ids = []
    for index, text in enumerate(["항상 보임", "1클릭", "2클릭"]):
        box = first.shapes.add_textbox(Inches(1), Inches(1 + index * 1.5), Inches(11), Inches(1.2))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        shape_ids.append(box.shape_id)

    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    lyrics = second.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(5))
    for index, line in enumerate(["첫째 줄", "둘째 줄", "셋째 줄"]):
        paragraph = lyrics.text_frame.paragraphs[0] if index == 0 else lyrics.text_frame.add_paragraph()
        paragraph.text = line
        paragraph.runs[0].font.size = Pt(40)
    lyrics_id = lyrics.shape_id

    third = presentation.slides.add_slide(presentation.slide_layouts[6])
    exit_ids = []
    for index, text in enumerate(["먼저 있다가 사라짐", "나중에 등장"]):
        box = third.shapes.add_textbox(Inches(1), Inches(1 + index * 2), Inches(11), Inches(1.2))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        exit_ids.append(box.shape_id)

    presentation.save(path)

    timings = {
        "ppt/slides/slide1.xml": _timing([
            _effect(10, shape_ids[1]),
            _effect(20, shape_ids[2]),
        ]),
        "ppt/slides/slide2.xml": _timing([
            _effect(30, lyrics_id, paragraph=1),
            _effect(40, lyrics_id, paragraph=2),
        ]),
        "ppt/slides/slide3.xml": _timing([
            _effect(50, exit_ids[1]),
            _effect(60, exit_ids[0], preset="exit"),
        ]),
    }

    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for name, timing in timings.items():
        root = etree.fromstring(parts[name])
        root.append(etree.fromstring(timing))
        parts[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    return shape_ids, lyrics_id, exit_ids


def slide_texts(path):
    """펼쳐진 pptx를 슬라이드 순서대로 [[문단 글자, ...], ...] 로 읽는다."""
    presentation = Presentation(str(path))
    pages = []
    for slide in presentation.slides:
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(p.text for p in shape.text_frame.paragraphs)
        pages.append(lines)
    return pages


def main():
    with tempfile.TemporaryDirectory(prefix="anim_test_") as work:
        work = Path(work)
        source = work / "sample.pptx"
        build_sample(source)

        assert len(slide_texts(source)) == 3, "원본은 3장이어야 한다"

        expanded = expand_animation_steps(source, work / "expanded.pptx")
        assert expanded is not None, "애니메이션을 못 찾았다"

        pages = slide_texts(expanded)
        assert len(pages) == 9, f"3장 × 3단계 = 9장이어야 하는데 {len(pages)}장"

        # 1장: 도형 단위 — 감춘 도형은 아예 사라진다.
        assert pages[0] == ["항상 보임"], pages[0]
        assert pages[1] == ["항상 보임", "1클릭"], pages[1]
        assert pages[2] == ["항상 보임", "1클릭", "2클릭"], pages[2]

        # 2장: 문단 단위 — 글자만 지우고 빈 줄은 남겨야 나머지 줄이 밀리지 않는다.
        assert pages[3] == ["첫째 줄", "", ""], pages[3]
        assert pages[4] == ["첫째 줄", "둘째 줄", ""], pages[4]
        assert pages[5] == ["첫째 줄", "둘째 줄", "셋째 줄"], pages[5]

        # 3장: 등장한 뒤 사라지는 것까지 순서대로 따라가야 한다.
        assert pages[6] == ["먼저 있다가 사라짐"], pages[6]
        assert pages[7] == ["먼저 있다가 사라짐", "나중에 등장"], pages[7]
        assert pages[8] == ["나중에 등장"], pages[8]

        # 애니메이션이 없으면 원본을 그대로 쓰라고 None을 돌려줘야 한다.
        plain = work / "plain.pptx"
        plain_presentation = Presentation()
        plain_presentation.slides.add_slide(plain_presentation.slide_layouts[6])
        plain_presentation.save(plain)
        assert expand_animation_steps(plain, work / "plain_expanded.pptx") is None

        print("ok: 3장 → 9장 (도형 단위 / 문단 단위 / 등장+사라짐 각 3단계)")


if __name__ == "__main__":
    main()

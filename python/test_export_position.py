"""가사 세로 미세 조정 self-check: python3 python/test_export_position.py

LibreOffice 없이 도는 단위 테스트. 상단/중단/하단 세 고정값만으로는 못 맞추는
높이를 `text_offset_y`(인치, + = 아래로)로 밀 수 있는지 본다.

규칙: 상단 여백(text_box_top)은 본문 상자의 "높이"를 정하고, 미세 조정은 그
상자를 통째로 위아래로 민다(높이는 그대로). 그래야 상단/중단/하단 어느 기준을
골라도 밀어 준 만큼 똑같이 움직인다.
"""
import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.util import Inches

from ppt_tool import (
    _TEXT_OFFSET_Y_MAX,
    _TEXT_OFFSET_Y_MIN,
    _lyrics_box_vertical_layout,
    export_presentation,
)

_STYLE = {
    "font_size": 40,
    "bible_font_size": 30,
    "text_box_top": 0.6,
    "bible_text_box_top": 0.6,
    "background_color": "#1B1B1B",
    "text_color": "#FFFFFF",
    "bible_text_color": "#FFF8E1",
    "text_position": "middle",
    "bible_text_position": "middle",
    "lyrics_text_align": "center",
    "bible_text_align": "center",
    "include_english_lyrics": False,
    "english_text_color": "#FFF176",
    "show_song_title": False,
    "show_bible_title": False,
    "title_font_size": 14,
    "bible_title_font_size": 14,
    "title_text_color": "#FFFFFF",
    "bible_title_text_color": "#FFFFFF",
    "title_horizontal_position": "right",
    "title_vertical_position": "bottom",
    "bible_title_horizontal_position": "right",
    "bible_title_vertical_position": "bottom",
    "font_family": "Pretendard",
    "background_image_path": None,
}

_failures = []


def check(label, actual, expected):
    if actual == expected:
        print(f"  OK   {label}")
    else:
        print(f"  FAIL {label}: {actual!r} != {expected!r}")
        _failures.append(label)


def close(label, actual, expected, tol=1e-6):
    if abs(actual - expected) <= tol:
        print(f"  OK   {label}")
    else:
        print(f"  FAIL {label}: {actual!r} != {expected!r}")
        _failures.append(label)


def body_box(slide):
    """가사 텍스트 상자. 제목을 끈 상태라 첫 도형이 본문이다."""
    for shape in slide.shapes:
        if shape.has_text_frame:
            return shape
    raise AssertionError("본문 상자가 없다")


def export(songs, style=None):
    tmp = tempfile.mkdtemp()
    out = Path(tmp) / "out.pptx"
    export_presentation(json.dumps({
        "output_path": str(out),
        "style": style or _STYLE,
        "songs": songs,
    }))
    return Presentation(str(out))


def test_layout_shifts_box_only():
    print("_lyrics_box_vertical_layout")
    base_top, base_height = _lyrics_box_vertical_layout(_STYLE, False)
    close("기본값은 그대로", base_top, 0.6)

    down_top, down_height = _lyrics_box_vertical_layout(
        dict(_STYLE, text_offset_y=0.75), False
    )
    close("아래로 0.75", down_top, base_top + 0.75)
    close("높이는 그대로", down_height, base_height)

    up_top, up_height = _lyrics_box_vertical_layout(
        dict(_STYLE, text_offset_y=-0.4), False
    )
    close("위로 0.4", up_top, base_top - 0.4)
    close("높이는 그대로 (위)", up_height, base_height)


def test_bible_uses_own_offset():
    print("찬양/성경 각각")
    style = dict(_STYLE, text_offset_y=0.5, bible_text_offset_y=-0.5)
    song_top, _ = _lyrics_box_vertical_layout(style, False)
    bible_top, _ = _lyrics_box_vertical_layout(style, True)
    close("가사는 아래로", song_top, 1.1)
    close("성경은 위로", bible_top, 0.1)


def test_offset_is_clamped():
    print("범위 밖 값")
    too_low, _ = _lyrics_box_vertical_layout(
        dict(_STYLE, text_offset_y=-99), False
    )
    too_high, _ = _lyrics_box_vertical_layout(
        dict(_STYLE, text_offset_y=99), False
    )
    close("최소로 잘림", too_low, 0.6 + _TEXT_OFFSET_Y_MIN)
    close("최대로 잘림", too_high, 0.6 + _TEXT_OFFSET_Y_MAX)

    junk, _ = _lyrics_box_vertical_layout(
        dict(_STYLE, text_offset_y="위로"), False
    )
    close("숫자가 아니면 0", junk, 0.6)

    nan, _ = _lyrics_box_vertical_layout(
        dict(_STYLE, text_offset_y=float("nan")), False
    )
    close("NaN 이면 0", nan, 0.6)


def test_export_moves_textbox():
    print("export (실제 상자 위치)")
    song = {
        "type": "song", "title": "찬양",
        "lyrics": "첫째 줄\n둘째 줄",
        "english_lyrics": "",
        "background": None,
    }
    base = body_box(export([song]).slides[0])
    moved = body_box(
        export([song], style=dict(_STYLE, text_offset_y=0.3)).slides[0]
    )
    # Inches() 변환에서 EMU 1 정도는 어긋난다 (0.3in = 274320 EMU)
    close("상자가 0.3인치 내려간다", moved.top - base.top, Inches(0.3), tol=2)
    check("상자 높이는 그대로", moved.height, base.height)
    check("가로는 그대로", (moved.left, moved.width), (base.left, base.width))


def test_export_without_key():
    print("export (하위 호환)")
    legacy = dict(_STYLE)
    legacy.pop("text_offset_y", None)
    prs = export([{
        "type": "song", "title": "찬양",
        "lyrics": "한 줄",
        "english_lyrics": "",
        "background": None,
    }], style=legacy)
    check("키가 없으면 기본 위치", body_box(prs.slides[0]).top, Inches(0.6))


if __name__ == "__main__":
    test_layout_shifts_box_only()
    test_bible_uses_own_offset()
    test_offset_is_clamped()
    test_export_moves_textbox()
    test_export_without_key()
    if _failures:
        print(f"\n{len(_failures)}개 실패: {_failures}")
        raise SystemExit(1)
    print("\n모두 통과")

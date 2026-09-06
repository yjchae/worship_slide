"""가사 줄바꿈 self-check: python3 python/test_export_text.py

LibreOffice 없이 도는 XML 단위 테스트. 여러 줄 가사가 <a:t> 안의 날 줄바꿈이 아니라
<a:br/> 로 들어가는지 본다. 날 줄바꿈으로 두면 OOXML 상 줄바꿈이 아니라서 뷰어가
가운데 정렬을 무시하고 줄을 양쪽으로 벌려 버린다.
"""
import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.oxml.ns import qn

from ppt_tool import export_presentation

_STYLE = {
    "font_size": 40,
    "bible_font_size": 30,
    "background_color": "#1B1B1B",
    "text_color": "#FFFFFF",
    "bible_text_color": "#FFF8E1",
    "text_position": "middle",
    "bible_text_position": "middle",
    "lyrics_text_align": "center",
    "bible_text_align": "center",
    "include_english_lyrics": True,
    "english_text_color": "#FFF176",
    "show_song_title": False,
    "show_bible_title": False,
    "title_font_size": 14,
    "bible_title_font_size": 14,
    "title_text_color": "#FFFFFF",
    "bible_title_text_color": "#FFFFFF",
    "title_horizontal_position": "right",
    "title_vertical_position": "bottom",
    "bible_title_horizontal_position": "right",
    "bible_title_vertical_position": "bottom",
    "font_family": "Pretendard",
    "background_image_path": None,
}

_failures = []


def check(label, actual, expected):
    if actual == expected:
        print(f"  OK   {label}")
    else:
        print(f"  FAIL {label}: {actual!r} != {expected!r}")
        _failures.append(label)


def body_paragraphs(slide):
    """가사 텍스트 상자의 문단들. 제목 상자는 별도 도형이라 첫 상자만 본다."""
    for shape in slide.shapes:
        if shape.has_text_frame:
            return shape.text_frame.paragraphs
    return []


def para_parts(paragraph):
    """문단 안 요소를 ('run', 텍스트) / ('br', None) 순서대로."""
    parts = []
    for child in paragraph._p:
        if child.tag == qn("a:r"):
            t = child.find(qn("a:t"))
            parts.append(("run", t.text if t is not None else ""))
        elif child.tag == qn("a:br"):
            parts.append(("br", None))
    return parts


def export(songs, style=None):
    tmp = tempfile.mkdtemp()
    out = Path(tmp) / "out.pptx"
    export_presentation(json.dumps({
        "output_path": str(out),
        "style": style or _STYLE,
        "songs": songs,
    }))
    return Presentation(str(out))


def test_line_breaks():
    print("가사 줄바꿈")
    prs = export([{
        "type": "song", "title": "찬양",
        "lyrics": "첫째 줄\n둘째 줄\n셋째 줄",
        "english_lyrics": "",
        "background": None,
    }])
    paragraphs = body_paragraphs(prs.slides[0])
    parts = para_parts(paragraphs[0])

    check("run / br 가 번갈아 들어간다", parts, [
        ("run", "첫째 줄"), ("br", None),
        ("run", "둘째 줄"), ("br", None),
        ("run", "셋째 줄"),
    ])
    # 어떤 <a:t> 에도 날 줄바꿈이 남으면 안 된다.
    raw_newlines = [text for kind, text in parts if kind == "run" and "\n" in (text or "")]
    check("a:t 안에 날 줄바꿈이 없다", raw_newlines, [])


def test_line_break_keeps_font():
    print("줄바꿈의 글꼴·크기")
    prs = export([{
        "type": "song", "title": "찬양",
        "lyrics": "첫째 줄\n둘째 줄",
        "english_lyrics": "",
        "background": None,
    }])
    paragraph = body_paragraphs(prs.slides[0])[0]
    br = paragraph._p.find(qn("a:br"))
    check("a:br 이 있다", br is not None, True)
    if br is None:
        return
    rpr = br.find(qn("a:rPr"))
    check("a:br 에 rPr 이 있다", rpr is not None, True)
    if rpr is None:
        return
    # 줄 높이가 본문과 같아야 한다. sz 는 100분의 1 pt.
    check("줄바꿈 크기 = 본문 크기", rpr.get("sz"), "4000")
    check("줄바꿈도 굵게", rpr.get("b"), "1")
    latin = rpr.find(qn("a:latin"))
    check("줄바꿈 글꼴", latin is not None and latin.get("typeface"), "Pretendard")
    # 스키마상 fill 이 latin/ea/cs 보다 앞이어야 한다.
    tags = [child.tag for child in rpr]
    check(
        "rPr 자식 순서 (solidFill → latin/ea/cs)",
        tags,
        [qn("a:solidFill"), qn("a:latin"), qn("a:ea"), qn("a:cs")],
    )


def test_single_line_has_no_break():
    print("한 줄짜리")
    prs = export([{
        "type": "song", "title": "찬양",
        "lyrics": "한 줄뿐",
        "english_lyrics": "",
        "background": None,
    }])
    parts = para_parts(body_paragraphs(prs.slides[0])[0])
    check("br 없이 run 하나", parts, [("run", "한 줄뿐")])


def test_english_lines():
    print("영어 가사")
    prs = export([{
        "type": "song", "title": "찬양",
        "lyrics": "한글 첫 줄\n한글 둘째 줄",
        "english_lyrics": "English one\nEnglish two",
        "background": None,
    }])
    paragraphs = body_paragraphs(prs.slides[0])
    check("한글 문단 + 영어 문단", len(paragraphs), 2)
    check("영어도 br 로 나뉜다", para_parts(paragraphs[1]), [
        ("run", "English one"), ("br", None), ("run", "English two"),
    ])


def test_alignment_kept():
    print("정렬 유지")
    for align, expected in (("center", "ctr"), ("left", "l"), ("right", "r")):
        prs = export(
            [{
                "type": "song", "title": "찬양",
                "lyrics": "첫째 줄\n둘째 줄",
                "english_lyrics": "",
                "background": None,
            }],
            style=dict(_STYLE, lyrics_text_align=align),
        )
        paragraph = body_paragraphs(prs.slides[0])[0]
        ppr = paragraph._p.find(qn("a:pPr"))
        actual = ppr.get("algn") if ppr is not None else None
        check(f"{align} 정렬", actual, expected)


def test_bible_unchanged():
    print("성경 본문 (줄마다 문단)")
    prs = export([{
        "type": "bible", "title": "롬 8:28",
        "lyrics": "첫째 줄\n둘째 줄\n셋째 줄",
        "english_lyrics": "",
        "background": None,
    }])
    paragraphs = body_paragraphs(prs.slides[0])
    # 성경은 원래 줄마다 문단을 만든다(내어쓰기 때문). br 이 끼면 안 된다.
    check("줄 수만큼 문단", len(paragraphs), 3)
    check("문단 안에는 br 이 없다", para_parts(paragraphs[0]), [("run", "첫째 줄")])


if __name__ == "__main__":
    test_line_breaks()
    test_line_break_keeps_font()
    test_single_line_has_no_break()
    test_english_lines()
    test_alignment_kept()
    test_bible_unchanged()
    if _failures:
        print(f"\n{len(_failures)}개 실패: {_failures}")
        raise SystemExit(1)
    print("\n모두 통과")

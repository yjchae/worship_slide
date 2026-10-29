"""추가 보조 언어(셋째 줄) self-check: python3 python/test_export_extra.py

LibreOffice 없이 도는 XML 단위 테스트. 원본 → 보조 → 추가 보조 순서로 문단이 쌓이는지,
색·크기(본문 × 비율, 성경은 성경 본문 기준)가 미리보기(Dart)와 같은 규칙인지 본다.
"""
import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.oxml.ns import qn

from ppt_tool import export_presentation

_STYLE = {
    "font_size": 40,
    "bible_font_size": 30,
    "background_color": "#1B1B1B",
    "text_color": "#FFFFFF",
    "bible_text_color": "#FFFFFF",
    "text_position": "middle",
    "bible_text_position": "middle",
    "lyrics_text_align": "center",
    "bible_text_align": "left",
    "include_english_lyrics": True,
    "english_text_color": "#FFF176",
    "include_extra_lyrics": True,
    "extra_text_color": "#80DEEA",
    "show_song_title": False,
    "show_bible_title": False,
    "title_font_size": 14,
    "bible_title_font_size": 14,
    "title_text_color": "#FFFFFF",
    "bible_title_text_color": "#FFFFFF",
    "title_horizontal_position": "right",
    "title_vertical_position": "bottom",
    "bible_title_horizontal_position": "right",
    "bible_title_vertical_position": "bottom",
    "font_family": "Pretendard",
    "background_image_path": None,
}

_failures = []


def check(label, actual, expected):
    if actual == expected:
        print(f"  OK   {label}")
    else:
        print(f"  FAIL {label}: {actual!r} != {expected!r}")
        _failures.append(label)


def export(songs, style=None):
    tmp = tempfile.mkdtemp()
    out = Path(tmp) / "out.pptx"
    export_presentation(json.dumps({
        "output_path": str(out),
        "style": style or _STYLE,
        "songs": songs,
    }))
    return Presentation(str(out))


def body_paragraphs(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            return shape.text_frame.paragraphs
    return []


def run_info(paragraph):
    """문단 첫 run 의 (글자, 크기 pt, 색 hex)."""
    run = paragraph._p.find(qn("a:r"))
    rpr = run.find(qn("a:rPr"))
    fill = rpr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    return run.find(qn("a:t")).text, int(rpr.get("sz")) / 100, fill.get("val")


def song(lyrics, english="", extra="", **extra_keys):
    return {
        "type": "song", "title": "찬양", "lyrics": lyrics,
        "english_lyrics": english, "extra_lyrics": extra,
        "background": None, **extra_keys,
    }


def test_three_languages():
    print("세 언어 문단")
    prs = export([song("한글", "English", "Español")])
    paragraphs = body_paragraphs(prs.slides[0])
    check("문단 세 개", len(paragraphs), 3)
    check("원본", run_info(paragraphs[0]), ("한글", 40, "FFFFFF"))
    check("보조 = 본문 × 0.8", run_info(paragraphs[1]), ("English", 32, "FFF176"))
    check("추가 보조 = 본문 × 0.7", run_info(paragraphs[2]), ("Español", 28, "80DEEA"))


def test_ratios_and_switch():
    print("비율 · 끄기")
    style = dict(_STYLE, sub_font_ratio=0.5, extra_font_ratio=0.25)
    paragraphs = body_paragraphs(export([song("한", "E", "X")], style).slides[0])
    check("보조 비율", run_info(paragraphs[1])[1], 20)
    check("추가 비율", run_info(paragraphs[2])[1], 10)

    off = dict(_STYLE, include_extra_lyrics=False)
    paragraphs = body_paragraphs(export([song("한", "E", "X")], off).slides[0])
    check("끄면 추가 보조 문단이 없다", len(paragraphs), 2)


def test_page_alignment():
    print("페이지 짝")
    prs = export([song("한1\n\n한2", "E1\n\nE2", "X1\n\nX2\n\nX3")])
    check("추가 보조만 있는 페이지도 한 장", len(prs.slides), 3)
    last = body_paragraphs(prs.slides[2])
    check("셋째 장은 추가 보조만", [p.text for p in last if p.text], ["X3"])


def test_bible_uses_bible_size():
    print("성경 크기")
    prs = export([{
        "type": "bible", "title": "요 3:16",
        "lyrics": "16. 하나님이", "english_lyrics": "16. For God",
        "extra_lyrics": "16. Porque", "background": None,
    }])
    paragraphs = body_paragraphs(prs.slides[0])
    texts = [run_info(p) for p in paragraphs if p._p.find(qn("a:r")) is not None]
    check("보조 = 성경 본문 30 × 0.8", texts[1][1], 24)
    check("추가 = 성경 본문 30 × 0.7", texts[2][1], 21)


def test_sermon_subtitle_size_wins():
    print("설교 부제 크기")
    prs = export([{
        "type": "blank", "title": "", "lyrics": "주제", "english_lyrics": "부제",
        "extra_lyrics": "", "background": {"sermon": True, "english_font_size": 50},
    }])
    paragraphs = body_paragraphs(prs.slides[0])
    check("부제는 항목이 정한 크기", run_info(paragraphs[1])[1], 50)


def test_blank_with_only_extra():
    print("빈 페이지 문구")
    prs = export([{
        "type": "blank", "title": "", "lyrics": "", "english_lyrics": "",
        "extra_lyrics": "Bienvenidos", "background": None,
    }])
    check("추가 보조 문구만 있어도 글 슬라이드", [p.text for p in body_paragraphs(prs.slides[0]) if p.text], ["Bienvenidos"])


if __name__ == "__main__":
    test_three_languages()
    test_ratios_and_switch()
    test_page_alignment()
    test_bible_uses_bible_size()
    test_sermon_subtitle_size_wins()
    test_blank_with_only_extra()
    if _failures:
        print(f"\n{len(_failures)}개 실패")
        sys.exit(1)
    print("\n모두 통과")

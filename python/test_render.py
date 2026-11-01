"""render 명령 self-check: python3 python/test_render.py [파일]

인자를 주지 않으면 tmp_sample/에서 첫 pptx를 집는다.
LibreOffice가 없으면 skip.
"""
import os
import sys
import tempfile
from pathlib import Path
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation

import ppt_tool
from ppt_tool import get_libreoffice_executable, render_presentation_images


def pick_sample():
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    sample_dir = Path(__file__).resolve().parent.parent / "tmp_sample"
    candidates = sorted(sample_dir.rglob("*.pptx"))
    return candidates[0] if candidates else None


def check_powerpoint_fails_without_libreoffice():
    """PowerPoint 가 못 구웠는데 LibreOffice 도 없으면 traceback 이 아니라 설치 안내용 오류."""
    with tempfile.TemporaryDirectory() as tmp:
        sample = Path(tmp) / "a.pptx"
        Presentation().save(str(sample))
        with (
            mock.patch.object(ppt_tool, "get_app_support_root", return_value=Path(tmp) / "cache"),
            mock.patch.object(ppt_tool, "get_libreoffice_executable", return_value=None),
            mock.patch.object(ppt_tool, "powerpoint_available", return_value=True),
            mock.patch.object(ppt_tool, "_convert_with_powerpoint", side_effect=RuntimeError("실패")),
        ):
            assert render_presentation_images(sample) == {"error": "libreoffice_missing"}


def main():
    check_powerpoint_fails_without_libreoffice()

    if get_libreoffice_executable() is None:
        print("skip: LibreOffice 없음")
        return

    sample = pick_sample()
    if sample is None:
        print("skip: 샘플 pptx 없음")
        return

    expected = len(Presentation(str(sample)).slides)
    result = render_presentation_images(sample)

    assert "error" not in result, result
    assert result["page_count"] == expected, (result["page_count"], expected)
    assert len(result["image_paths"]) == expected
    for path in result["image_paths"]:
        assert os.path.getsize(path) > 0, path

    # 두 번째 호출은 캐시에서 같은 결과가 나와야 한다.
    assert render_presentation_images(sample)["image_paths"] == result["image_paths"]

    print(f"ok: {sample.name} → {expected}장")


if __name__ == "__main__":
    main()

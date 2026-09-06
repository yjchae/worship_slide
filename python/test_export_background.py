"""항목별 배경 오버라이드 self-check: python3 python/test_export_background.py

LibreOffice 없이 도는 XML 단위 테스트. 콘티 일부 항목만 배경을 다르게 지정했을 때
export가 만든 pptx의 슬라이드 배경이 항목별로 갈리는지 확인한다.
(헌금송만 다른 배경 — 나머지는 전역 배경 그대로)
"""
import base64
import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.oxml.ns import qn

from ppt_tool import _resolve_background, export_presentation

# 1x1 투명 PNG. python-pptx가 크기를 읽을 수 있는 최소 이미지.
_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)

GLOBAL_BG = "#1B1B1B"
OFFERING_BG = "#123456"
BLANK_BG = "#654321"

_STYLE = {
    "font_size": 30,
    "bible_font_size": 30,
    "text_box_top": 0.6,
    "bible_text_box_top": 0.6,
    "background_color": GLOBAL_BG,
    "text_color": "#FFFFFF",
    "bible_text_color": "#FFFFFF",
    "text_position": "middle",
    "bible_text_position": "middle",
    "lyrics_text_align": "center",
    "bible_text_align": "center",
    "include_english_lyrics": False,
    "english_text_color": "#FFF176",
    "show_song_title": False,
    "show_bible_title": False,
    "title_font_size": 14,
    "bible_title_font_size": 14,
    "title_text_color": "#FFFFFF",
    "bible_title_text_color": "#FFFFFF",
    "title_horizontal_position": "right",
    "title_vertical_position": "bottom",
    "bible_title_horizontal_position": "right",
    "bible_title_vertical_position": "bottom",
    "font_family": "Pretendard",
    "background_image_path": None,
}

_failures = []


def check(label, actual, expected):
    if actual == expected:
        print(f"  OK   {label}")
    else:
        print(f"  FAIL {label}: {actual!r} != {expected!r}")
        _failures.append(label)


def _bg_element(slide):
    # p:bg 는 p:sld 바로 밑이 아니라 p:cSld 안에 있다.
    return slide._element.find(f'{qn("p:cSld")}/{qn("p:bg")}')


def slide_bg_hex(slide):
    """슬라이드 배경의 단색 fill 을 '#RRGGBB' 로. 이미지 배경이면 None."""
    bg = _bg_element(slide)
    if bg is None or bg.find(f'.//{qn("a:blipFill")}') is not None:
        return None
    return "#" + str(slide.background.fill.fore_color.rgb)


def slide_bg_is_image(slide):
    bg = _bg_element(slide)
    return bg is not None and bg.find(f'.//{qn("a:blipFill")}') is not None


def test_resolve_background():
    print("_resolve_background")
    # 오버라이드가 없으면 전역 배경 그대로
    check(
        "오버라이드 없음 → 전역",
        _resolve_background(_STYLE, None),
        (None, GLOBAL_BG),
    )
    # 색만 담긴 오버라이드는 "이 항목만 단색" (전역 배경 이미지를 덮는다)
    style_with_image = dict(_STYLE, background_image_path="/tmp/global.png")
    check(
        "색 오버라이드가 전역 배경 이미지를 덮는다",
        _resolve_background(style_with_image, {"color": OFFERING_BG}),
        (None, OFFERING_BG),
    )
    # 이미지까지 담긴 오버라이드
    check(
        "이미지 오버라이드",
        _resolve_background(
            _STYLE, {"color": OFFERING_BG, "image_path": "/tmp/a.png"}
        ),
        ("/tmp/a.png", OFFERING_BG),
    )
    # 색이 비어 있으면 전역 색으로 떨어진다 (손상된 값 방어)
    check(
        "색 없는 오버라이드 → 전역 색",
        _resolve_background(_STYLE, {"image_path": "/tmp/a.png"}),
        ("/tmp/a.png", GLOBAL_BG),
    )


def test_export_per_item_background():
    print("export (항목별 배경)")
    with tempfile.TemporaryDirectory() as tmp:
        image_path = Path(tmp) / "offering_bg.png"
        image_path.write_bytes(_PNG_1X1)
        output = Path(tmp) / "out.pptx"

        payload = {
            "output_path": str(output),
            "style": _STYLE,
            "songs": [
                {
                    "type": "song",
                    "title": "헌금송",
                    "lyrics": "헌금송 1절",
                    "english_lyrics": "",
                    "background": {"color": OFFERING_BG},
                },
                {
                    "type": "song",
                    "title": "일반 찬양",
                    "lyrics": "일반 1절",
                    "english_lyrics": "",
                    "background": None,
                },
                {
                    "type": "blank",
                    "title": "",
                    "lyrics": "",
                    "english_lyrics": "",
                    "background": {"color": BLANK_BG},
                },
                {
                    "type": "song",
                    "title": "이미지 배경 찬양",
                    "lyrics": "이미지 1절",
                    "english_lyrics": "",
                    "background": {
                        "color": OFFERING_BG,
                        "image_path": str(image_path),
                    },
                },
            ],
        }

        export_presentation(json.dumps(payload))
        slides = Presentation(str(output)).slides

        # 헌금송 1장 + 뒤따르는 자동 여백 1장 + 일반 찬양 1장
        # + 명시적 빈 페이지 1장 + 이미지 배경 찬양 1장
        check("슬라이드 개수", len(slides), 5)
        check("헌금송 배경", slide_bg_hex(slides[0]), OFFERING_BG)
        check("헌금송 뒤 자동 여백도 같은 배경", slide_bg_hex(slides[1]), OFFERING_BG)
        check("오버라이드 없는 찬양은 전역 배경", slide_bg_hex(slides[2]), GLOBAL_BG)
        check("빈 페이지 오버라이드", slide_bg_hex(slides[3]), BLANK_BG)
        check("이미지 오버라이드는 blipFill", slide_bg_is_image(slides[4]), True)


def test_export_without_background_key():
    """예전 payload(배경 키 없음)도 그대로 돌아야 한다."""
    print("export (하위 호환)")
    with tempfile.TemporaryDirectory() as tmp:
        output = Path(tmp) / "legacy.pptx"
        payload = {
            "output_path": str(output),
            "style": _STYLE,
            "songs": [
                {
                    "type": "song",
                    "title": "찬양",
                    "lyrics": "1절",
                    "english_lyrics": "",
                }
            ],
        }
        export_presentation(json.dumps(payload))
        slides = Presentation(str(output)).slides
        check("배경 키가 없으면 전역 배경", slide_bg_hex(slides[0]), GLOBAL_BG)


if __name__ == "__main__":
    test_resolve_background()
    test_export_per_item_background()
    test_export_without_background_key()
    if _failures:
        print(f"\n{len(_failures)}개 실패: {_failures}")
        raise SystemExit(1)
    print("\n모두 통과")
